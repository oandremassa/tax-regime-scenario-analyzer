from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

try:
    from openpyxl import load_workbook
except Exception:  # optional at import time in minimal environments
    load_workbook = None


MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
MONEY_RE = re.compile(r"-?\d[\d\.]*,\d{2}|-?\d+(?:\.\d+)?")


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _num(value):
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _normalize_header(value):
    s = str(value or "").strip().lower()
    repl = str.maketrans("áàãâéêíóôõúç", "aaaaeeiooouc")
    return s.translate(repl).replace(" ", "_").replace("-", "_")


def _monthly_from_rows(headers, rows):
    aliases = {
        "month": {"month", "mes", "competencia", "period", "periodo"},
        "commerce_revenue": {"commerce_revenue", "receita_comercio", "comercio", "commerce"},
        "industry_revenue": {"industry_revenue", "receita_industria", "industria", "industry"},
        "service_revenue": {"service_revenue", "receita_servicos", "servicos", "services", "service"},
        "payroll": {"payroll", "folha", "folha_pagamento", "salary", "salarios"},
        "costs": {"costs", "custos", "cost"},
        "expenses": {"expenses", "despesas", "expense"},
        "current_tax_paid": {"current_tax_paid", "tax_paid", "imposto_pago", "das", "current_tax"},
    }
    normalized = [_normalize_header(x) for x in headers]
    index = {}
    for target, options in aliases.items():
        for i, h in enumerate(normalized):
            if h in options:
                index[target] = i
                break

    if "month" not in index:
        return []
    revenue_keys = {"commerce_revenue", "industry_revenue", "service_revenue"}
    if not revenue_keys.intersection(index):
        # generic revenue is accepted as service revenue for import convenience
        for i, h in enumerate(normalized):
            if h in {"revenue", "receita", "faturamento", "gross_revenue"}:
                index["service_revenue"] = i
                break
    if not revenue_keys.intersection(index):
        return []

    result = []
    for row in rows:
        month_raw = row[index["month"]] if index["month"] < len(row) else None
        month = str(month_raw or "").strip()[:7]
        if not MONTH_RE.match(month):
            continue
        item = {"month": month}
        for key in revenue_keys | {"payroll", "costs", "expenses", "current_tax_paid"}:
            i = index.get(key)
            item[key] = _num(row[i]) if i is not None and i < len(row) else 0.0
        result.append(item)
    return result


def parse_csv(path: Path):
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = text[:3000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return {"parser": "csv", "doc_type": "csv", "warnings": ["Empty CSV file."]}
    monthly = _monthly_from_rows(rows[0], rows[1:])
    return {
        "parser": "csv_financials" if monthly else "csv_generic",
        "doc_type": "financial_csv" if monthly else "csv",
        "monthly_rows": monthly,
        "rows_detected": max(len(rows) - 1, 0),
        "warnings": [] if monthly else ["CSV stored, but no recognized monthly financial structure was detected."],
    }


def parse_xlsx(path: Path):
    if load_workbook is None:
        return {"parser": "xlsx", "doc_type": "spreadsheet", "warnings": ["Spreadsheet parser unavailable in this environment."]}
    wb = load_workbook(path, read_only=True, data_only=True)
    best = None
    for ws in wb.worksheets:
        data = []
        for row in ws.iter_rows(values_only=True):
            data.append(list(row))
            if len(data) > 500:
                break
        if not data:
            continue
        monthly = _monthly_from_rows(data[0], data[1:])
        candidate = {
            "parser": "xlsx_financials" if monthly else "xlsx_generic",
            "doc_type": "dre" if monthly else "spreadsheet",
            "sheet": ws.title,
            "monthly_rows": monthly,
            "rows_detected": max(len(data) - 1, 0),
            "warnings": [] if monthly else ["Spreadsheet stored; no recognized monthly structure was detected on the selected sheet."],
        }
        if monthly:
            best = candidate
            break
        if best is None:
            best = candidate
    return best or {"parser": "xlsx", "doc_type": "spreadsheet", "warnings": ["No readable worksheet found."]}


def _xml_local(tag):
    return tag.split("}")[-1] if "}" in tag else tag


def parse_xml(path: Path):
    try:
        root = ET.parse(path).getroot()
    except ET.ParseError as exc:
        return {"parser": "xml", "doc_type": "xml", "warnings": [f"Invalid XML: {exc}"]}

    values = {}
    for elem in root.iter():
        key = _xml_local(elem.tag)
        text = (elem.text or "").strip()
        if text and key not in values:
            values[key] = text
    amount = _num(values.get("vNF") or values.get("vServ") or values.get("valor"))
    cnpj = values.get("CNPJ")
    return {
        "parser": "nfe_xml_proxy" if "vNF" in values else "xml_generic",
        "doc_type": "nfe" if "vNF" in values else "xml",
        "document_identifier": values.get("nNF") or values.get("numero") or values.get("Id"),
        "cnpj_detected": cnpj,
        "amount_detected": amount,
        "issue_date": values.get("dhEmi") or values.get("dEmi"),
        "warnings": [] if values else ["XML has no readable fields."],
    }


def parse_txt(path: Path):
    text = path.read_text(encoding="utf-8", errors="replace")
    upper = text.upper()
    parsed = {"parser": "txt_generic", "doc_type": "txt", "warnings": []}
    if "PGDAS" in upper or "SIMPLES NACIONAL" in upper:
        parsed["parser"] = "pgdas_proxy"
        parsed["doc_type"] = "pgdas"
        patterns = {
            "rbt12": r"RBT12\s*[:=]\s*([\d\.,]+)",
            "current_tax": r"(?:DAS|TOTAL)\s*[:=]\s*([\d\.,]+)",
            "revenue_period": r"(?:RECEITA|FATURAMENTO)\s*[:=]\s*([\d\.,]+)",
            "cnpj_detected": r"CNPJ\s*[:=]\s*([\d\.\-/]+)",
        }
        for key, pattern in patterns.items():
            m = re.search(pattern, text, flags=re.I)
            if m:
                parsed[key] = m.group(1) if key == "cnpj_detected" else _num(m.group(1))
    return parsed


def parse_document(path: str | Path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return parse_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_xlsx(path)
    if suffix == ".xml":
        return parse_xml(path)
    if suffix in {".txt", ".dat"}:
        return parse_txt(path)
    if suffix == ".pdf":
        return {"parser": "pdf_stored", "doc_type": "pdf", "warnings": ["PDF stored for traceability; manual review is required in this portfolio build."]}
    return {"parser": "stored_only", "doc_type": suffix.lstrip(".") or "file", "warnings": ["File stored without structured parsing."]}
